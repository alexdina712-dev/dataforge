"""Export data as values; CSV formula-like text is neutralized deliberately."""

from io import BytesIO
import re
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from .profiling import scalar


def safe_csv_text(value):
    return "'" + value if isinstance(value, str) and re.match(r"^\s*[=+\-@\t\r]", value) else value


def export_csv(frame: pd.DataFrame) -> bytes:
    safe = frame.copy(deep=True)
    safe.columns = [safe_csv_text(c) for c in safe.columns]
    for column in safe:
        safe[column] = safe[column].map(lambda v: safe_csv_text(scalar(v)))
    return safe.to_csv(index=False, lineterminator="\r\n").encode("utf-8-sig")


def export_xlsx(frame: pd.DataFrame) -> bytes:
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.utils import get_column_letter

    book = Workbook(write_only=True)
    sheet = book.create_sheet("Cleaned data")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(frame.columns))}{len(frame) + 1}"
    for number, column in enumerate(frame.columns, 1):
        values = [str(scalar(v) or "") for v in frame[column].head(100)]
        sheet.column_dimensions[get_column_letter(number)].width = min(
            45, max(12, len(column) + 3, *(len(v) + 3 for v in values))
        )
    header = []
    for column in frame.columns:
        cell = WriteOnlyCell(sheet, value=column)
        cell.data_type = "s"
        cell.font = Font(name="Calibri", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="292F50")
        cell.alignment = Alignment(vertical="center")
        header.append(cell)
    sheet.append(header)
    for row in frame.itertuples(index=False, name=None):
        cells = []
        for value in row:
            cell = WriteOnlyCell(sheet, value=scalar(value))
            if isinstance(cell.value, str):
                cell.data_type = "s"
            cells.append(cell)
        sheet.append(cells)
    buffer = BytesIO()
    book.save(buffer)
    book.close()
    return buffer.getvalue()
