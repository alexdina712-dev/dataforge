"""Parse bounded CSV/OOXML content without pandas' permissive header/NA defaults."""

import csv
from datetime import date, datetime
from io import BytesIO, StringIO
import math
from pathlib import Path
import re
from zipfile import ZipFile
from defusedxml.ElementTree import iterparse
from openpyxl.utils.cell import column_index_from_string
import pandas as pd
from openpyxl import load_workbook
from .config import (
    MAX_UPLOAD_BYTES,
    MAX_ROWS,
    MAX_COLUMNS,
    MAX_CELLS,
    MAX_CELL_CHARS,
    MAX_SAFE_NUMBER,
)
from .errors import DataError


def validate_names(names: list[str]) -> None:
    if not names or len(names) > MAX_COLUMNS:
        raise DataError(f"Use between 1 and {MAX_COLUMNS} columns.")
    if any(
        not n.strip() or len(n) > 80 or any(ord(c) < 32 or ord(c) == 127 for c in n) for n in names
    ):
        raise DataError(
            "Column headers must be non-empty, at most 80 characters, and contain no control characters."
        )
    if len(names) != len(set(names)):
        raise DataError(
            "Duplicate column headers are not supported. Rename them in the source file first."
        )


def kind(series: pd.Series) -> str:
    dtype = series.dtype
    if pd.api.types.is_bool_dtype(dtype):
        return "boolean"
    if pd.api.types.is_datetime64_any_dtype(dtype):
        return "date"
    if pd.api.types.is_integer_dtype(dtype):
        return "integer"
    if pd.api.types.is_numeric_dtype(dtype):
        return "decimal"
    return "text"


def infer_frame(frame: pd.DataFrame) -> pd.DataFrame:
    result = frame.copy(deep=True)
    for column in result:
        s = result[column].map(
            lambda v: None if v is None or isinstance(v, str) and not v.strip() else v
        )
        strings = s.map(lambda v: None if pd.isna(v) else str(v))
        observed = strings.dropna()
        if observed.empty:
            result[column] = strings.astype("string")
            continue
        stripped = observed.str.strip()
        lowered = stripped.str.lower()
        if lowered.isin(["true", "false"]).all():
            result[column] = strings.map(
                lambda v: pd.NA if v is None else v.strip().lower() == "true"
            ).astype("boolean")
            continue
        if stripped.str.fullmatch(r"\d{4}-\d{2}-\d{2}").all():
            parsed = pd.to_datetime(strings.str.strip(), format="%Y-%m-%d", errors="coerce")
            if parsed.notna().sum() == len(observed):
                result[column] = parsed
                continue
        numeric = stripped.str.fullmatch(r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?").all()
        leading_zero = stripped.str.fullmatch(r"[+-]?0\d+").any()
        if numeric and not leading_zero:
            try:
                nums = pd.to_numeric(strings.str.strip(), errors="raise")
                valid = nums.dropna().astype(float)
                if valid.map(math.isfinite).all() and valid.abs().le(MAX_SAFE_NUMBER).all():
                    result[column] = nums.astype("Int64" if (valid % 1 == 0).all() else "Float64")
                    continue
            except (ValueError, TypeError, OverflowError):
                pass
        result[column] = strings.astype("string")
    return result


def _cell(value):
    if value is None:
        return None
    if isinstance(value, datetime):
        return (
            value.date().isoformat()
            if value.time().isoformat() == "00:00:00"
            else value.isoformat()
        )
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and not math.isfinite(value):
        raise DataError("The workbook contains a non-finite numeric value.")
    if isinstance(value, str) and re.search(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", value):
        raise DataError("Cell text contains unsupported control characters.")
    text = str(value)
    if len(text) > MAX_CELL_CHARS:
        raise DataError(f"Cell text exceeds the {MAX_CELL_CHARS}-character limit.")
    return None if isinstance(value, str) and not value.strip() else value


def _frame(names, rows):
    names = [str(n) if n is not None else "" for n in names]
    validate_names(names)
    if not rows:
        raise DataError("The dataset has headers but no data rows.")
    if len(rows) > MAX_ROWS or len(rows) * len(names) > MAX_CELLS:
        raise DataError(f"Dataset exceeds {MAX_ROWS:,} rows or {MAX_CELLS:,} cells.")
    return infer_frame(pd.DataFrame(rows, columns=names, dtype=object))


def inspect_workbook(data: bytes) -> list[str]:
    try:
        with ZipFile(BytesIO(data)) as archive:
            items = archive.infolist()
            if len(items) > 1000 or sum(i.file_size for i in items) > 20 * 1024 * 1024:
                raise DataError("Workbook archive exceeds safe expansion limits.")
            names = [i.filename.lower() for i in items]
            if "[content_types].xml" not in names or "xl/workbook.xml" not in names:
                raise DataError("This is not a valid XLSX workbook.")
            if any(
                "vbaproject" in n
                or "externallinks/" in n
                or "embeddings/" in n
                or n == "xl/connections.xml"
                for n in names
            ):
                raise DataError(
                    "Macros, embedded objects, external links, and data connections are not supported."
                )
            if any(
                i.file_size > 10 * 1024 * 1024 or i.file_size > max(1, i.compress_size) * 200
                for i in items
            ):
                raise DataError("Workbook archive compression exceeds safe limits.")
            cells = 0
            for item in items:
                if not item.filename.startswith("xl/worksheets/") or not item.filename.endswith(
                    ".xml"
                ):
                    continue
                with archive.open(item) as xml:
                    for event, element in iterparse(xml, events=("start", "end")):
                        tag = element.tag.rsplit("}", 1)[-1]
                        if event == "start" and tag == "f":
                            raise DataError(
                                "Formulas are not supported. Save worksheets as values before uploading."
                            )
                        if event == "start" and tag == "c":
                            ref = element.attrib.get("r", "")
                            match = re.fullmatch(r"([A-Z]{1,3})([1-9][0-9]{0,7})", ref)
                            if (
                                not match
                                or column_index_from_string(match[1]) > MAX_COLUMNS
                                or int(match[2]) > MAX_ROWS + 1
                            ):
                                raise DataError(
                                    "Worksheet cell coordinates exceed the row/column limits."
                                )
                            cells += 1
                            if cells > MAX_CELLS + MAX_COLUMNS * 20:
                                raise DataError("Workbook exceeds the cell limit.")
                        if event == "end":
                            element.clear()
        book = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
        try:
            if len(book.sheetnames) > 20:
                raise DataError("A workbook may contain at most 20 worksheets.")
            return book.sheetnames
        finally:
            book.close()
    except DataError:
        raise
    except Exception as exc:
        raise DataError("The XLSX workbook is malformed or unreadable.") from exc


def validate_file(filename: str, data: bytes) -> str:
    extension = Path(filename).suffix.lower()
    if extension not in [".csv", ".xlsx"]:
        raise DataError(
            "Upload a CSV or XLSX file. XLS, XLSM, macros, and other formats are not supported."
        )
    if not data:
        raise DataError("The uploaded file is empty.")
    if len(data) > MAX_UPLOAD_BYTES:
        raise DataError("Files must be 5 MB or smaller.", 413)
    return extension


def parse_file(
    filename: str, data: bytes, sheet: str | None = None, delimiter: str = "auto"
) -> tuple[pd.DataFrame, dict]:
    extension = validate_file(filename, data)
    if extension == ".csv":
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise DataError(
                "CSV files must use UTF-8 encoding. Export your file as UTF-8 CSV and try again."
            ) from exc
        if "\x00" in text:
            raise DataError("CSV files cannot contain binary or null bytes.")
        choices = {"comma": ",", "semicolon": ";", "tab": "\t", "pipe": "|"}
        if delimiter not in ["auto", *choices]:
            raise DataError("Choose a supported CSV delimiter.")
        if delimiter == "auto":
            try:
                separator = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|").delimiter
            except csv.Error:
                separator = ","
        else:
            separator = choices[delimiter]
        try:
            reader = csv.reader(StringIO(text, newline=""), delimiter=separator, strict=True)
            header = next((r for r in reader if r), None)
            if header is None:
                raise DataError("The CSV has no header row.")
            validate_names(header)
            rows = []
            for row in reader:
                if not row:
                    continue
                if len(row) != len(header):
                    raise DataError(
                        f"CSV row {reader.line_num} has {len(row)} fields; expected {len(header)}. Check the delimiter or repair the row."
                    )
                rows.append([_cell(v) for v in row])
                if len(rows) > MAX_ROWS or len(rows) * len(header) > MAX_CELLS:
                    raise DataError("Dataset exceeds the row/cell limit.")
        except csv.Error as exc:
            raise DataError("The CSV contains malformed quoting.") from exc
        return _frame(header, rows), {
            "format": "CSV",
            "delimiter": separator,
            "sheet": None,
            "sheets": [],
        }
    sheets = inspect_workbook(data)
    selected = sheet or sheets[0]
    if selected not in sheets:
        raise DataError("The selected worksheet does not exist.")
    book = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
    try:
        ws = book[selected]
        ws.reset_dimensions()
        header = None
        rows = []
        for cells in ws.iter_rows():
            if len(cells) > MAX_COLUMNS:
                raise DataError(f"Worksheet exceeds {MAX_COLUMNS} columns.")
            if any(c.data_type == "f" for c in cells):
                raise DataError(
                    "Formulas are not supported. Save the worksheet as values before uploading."
                )
            values = [_cell(c.value) for c in cells]
            if header is None:
                if not any(v is not None for v in values):
                    continue
                while values and values[-1] is None:
                    values.pop()
                header = values
                validate_names([str(v) if v is not None else "" for v in header])
                continue
            if len(values) > len(header) and any(v is not None for v in values[len(header) :]):
                raise DataError("Worksheet contains data outside the header columns.")
            rows.append((values[: len(header)] + [None] * len(header))[: len(header)])
            if len(rows) > MAX_ROWS or len(rows) * len(header) > MAX_CELLS:
                raise DataError("Worksheet exceeds the row/cell limit.")
        while rows and all(v is None for v in rows[-1]):
            rows.pop()
        if header is None:
            raise DataError("The worksheet is empty.")
        return _frame(header, rows), {
            "format": "XLSX",
            "delimiter": None,
            "sheet": selected,
            "sheets": sheets,
        }
    except DataError:
        raise
    except Exception as exc:
        raise DataError("The worksheet is malformed or unreadable.") from exc
    finally:
        book.close()
