from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
from pathlib import Path
import pandas as pd
from pandas.testing import assert_frame_equal
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from app.config import Settings
from app.errors import DataError
from app.main import create_app
from app.storage import WorkspaceStore
from app.schemas import Operation

DIRTY = (Path(__file__).parent / "fixtures/dirty.csv").read_bytes()


@pytest.fixture
def client():
    with TestClient(create_app(Settings())) as c:
        assert c.get("/api/session").status_code == 200
        yield c


def upload(client, data=DIRTY, name="dirty.csv"):
    response = client.post("/api/datasets", files={"file": (name, data, "text/csv")})
    assert response.status_code == 201, response.text
    return response.json()


def transform(client, dataset, operation, **kwargs):
    return client.post(
        "/api/datasets/" + dataset["id"] + "/transform",
        json={"revision": dataset["revision"], "transform": {"operation": operation, **kwargs}},
    )


def test_session_security_and_no_store_headers(client):
    assert client.get("/api/health").json()["status"] == "ok"
    assert client.get("/api/datasets").headers["cache-control"] == "no-store"
    assert client.cookies.get("dataforge_session")
    with TestClient(create_app(Settings())) as other:
        assert other.get("/api/datasets").status_code == 401


def test_upload_profiles_original_and_current(client):
    d = upload(client)
    assert d["profile"]["rows"] == 4
    assert d["profile"]["missingCells"] == 3
    assert d["profile"]["duplicateRows"] == 1
    assert d["originalPreview"] == d["currentPreview"]
    assert d["revision"] == 1
    assert d["history"] == []
    assert len(client.get("/api/datasets").json()) == 1


def test_samples_are_private_copies_and_workbook_inspection(client):
    samples = client.get("/api/samples").json()
    assert len(samples) == 3
    sales = client.post("/api/samples/sales")
    assert sales.status_code == 201, sales.text
    assert sales.json()["profile"]["rows"] == 27
    assert client.post("/api/samples/unknown").status_code == 404
    path = Path(__file__).resolve().parents[2] / "samples/inventory-review.xlsx"
    response = client.post("/api/files/inspect", files={"file": ("stock.xlsx", path.read_bytes())})
    assert response.status_code == 200
    assert response.json()["sheets"] == ["Stock", "Suppliers"]
    response = client.post(
        "/api/datasets",
        files={"file": ("stock.xlsx", path.read_bytes())},
        data={"sheet": "Suppliers"},
    )
    assert response.status_code == 201, response.text
    assert response.json()["profile"]["rows"] == 2


def test_preview_is_read_only_and_apply_changes_only_current(client):
    d = upload(client)
    body = {"revision": 1, "transform": {"operation": "remove_duplicates"}}
    dry = client.post("/api/datasets/" + d["id"] + "/preview-transform", json=body)
    assert dry.status_code == 200
    assert dry.json()["after"]["rows"] == 3
    assert client.get("/api/datasets/" + d["id"]).json()["revision"] == 1
    applied = client.post("/api/datasets/" + d["id"] + "/transform", json=body).json()
    assert applied["profile"]["rows"] == 3
    assert applied["originalProfile"]["rows"] == 4
    assert len(applied["history"]) == 1
    assert client.post("/api/datasets/" + d["id"] + "/transform", json=body).status_code == 409


def test_history_undo_replays_and_restores_exact_dtypes(client):
    d = upload(client)
    d = transform(client, d, "fill_mean", column="Revenue").json()
    d = transform(client, d, "normalize_names").json()
    original = client.app.state.store.frame(
        client.cookies["dataforge_session"], d["id"], "original"
    ).copy(deep=True)
    d = client.post("/api/datasets/" + d["id"] + "/undo", json={"revision": d["revision"]}).json()
    assert "Revenue" in d["currentPreview"]["columns"]
    assert d["profile"]["missingCells"] == 2
    d = client.post("/api/datasets/" + d["id"] + "/undo", json={"revision": d["revision"]}).json()
    assert d["history"] == []
    assert_frame_equal(
        client.app.state.store.frame(client.cookies["dataforge_session"], d["id"]), original
    )
    assert (
        client.post(
            "/api/datasets/" + d["id"] + "/undo", json={"revision": d["revision"]}
        ).status_code
        == 409
    )


def test_reset_restores_original_and_clears_history(client):
    d = upload(client)
    d = transform(client, d, "drop_columns", columns=["Region"]).json()
    assert d["profile"]["columns"] == 3
    r = client.post("/api/datasets/" + d["id"] + "/reset", json={"revision": d["revision"]})
    assert r.status_code == 200
    assert r.json()["profile"]["columns"] == 4
    assert r.json()["history"] == []


def test_private_ownership_all_data_endpoints(client):
    d = upload(client)
    with TestClient(client.app) as other:
        other.get("/api/session")
        for suffix in [
            "",
            "/preview",
            "/export?format=csv",
            "/export?format=xlsx",
            "/duplicates?column=Region",
            "/chart?kind=histogram&x=Revenue",
        ]:
            assert other.get("/api/datasets/" + d["id"] + suffix).status_code == 404
        for suffix, body in [
            ("/transform", {"revision": 1, "transform": {"operation": "remove_duplicates"}}),
            (
                "/preview-transform",
                {"revision": 1, "transform": {"operation": "remove_duplicates"}},
            ),
            ("/undo", {"revision": 1}),
            ("/reset", {"revision": 1}),
        ]:
            assert other.post("/api/datasets/" + d["id"] + suffix, json=body).status_code == 404
        assert other.delete("/api/datasets/" + d["id"]).status_code == 404
        assert other.get("/api/datasets").json() == []


def test_malformed_uploads_and_validation_are_safe(client):
    malformed = (Path(__file__).parent / "fixtures/malformed-row.csv").read_bytes()
    response = client.post("/api/datasets", files={"file": ("bad.csv", malformed)})
    assert response.status_code == 400
    assert "unexpected" not in response.text
    assert (
        client.post("/api/datasets", files={"file": ("script.py", b"print(1)")}).status_code == 400
    )
    assert (
        client.post(
            "/api/datasets", files={"file": ("large.csv", b"x" * (5 * 1024 * 1024 + 1))}
        ).status_code
        == 413
    )
    assert (
        client.post(
            "/api/datasets",
            content=b"x" * (6 * 1024 * 1024),
            headers={"Content-Type": "application/octet-stream"},
        ).status_code
        == 413
    )
    assert client.get("/api/datasets").json() == []


def test_arbitrary_code_and_unknown_fields_are_rejected(client):
    d = upload(client)
    for op in [
        {"operation": '__import__("os").system("echo no")'},
        {"operation": "remove_duplicates", "code": "print(1)"},
    ]:
        r = client.post(
            "/api/datasets/" + d["id"] + "/transform", json={"revision": 1, "transform": op}
        )
        assert r.status_code == 422
        assert "print(1)" not in r.text
        assert "__import__" not in r.text
    assert client.get("/api/datasets/" + d["id"]).json()["revision"] == 1


def test_chart_candidates_pagination_and_export(client):
    d = upload(client)
    assert (
        client.get("/api/datasets/" + d["id"] + "/preview?offset=1&limit=2").json()["rows"][0][0]
        == "002"
    )
    assert client.get("/api/datasets/" + d["id"] + "/preview?limit=500").status_code == 422
    chart = client.get("/api/datasets/" + d["id"] + "/chart?kind=histogram&x=Revenue")
    assert chart.status_code == 200
    assert chart.json()["includedRows"] == 3
    assert client.get("/api/datasets/" + d["id"] + "/duplicates?column=Region").status_code == 200
    csv = client.get("/api/datasets/" + d["id"] + "/export?format=csv")
    assert csv.status_code == 200
    assert csv.content.startswith(b"\xef\xbb\xbf")
    assert "attachment" in csv.headers["content-disposition"]
    excel = client.get("/api/datasets/" + d["id"] + "/export?format=xlsx")
    assert excel.status_code == 200
    book = load_workbook(BytesIO(excel.content))
    assert book.active.max_row == 5
    book.close()


def test_dataset_and_workspace_deletion(client):
    first = upload(client)
    assert client.delete("/api/datasets/" + first["id"]).status_code == 204
    assert client.get("/api/datasets/" + first["id"]).status_code == 404
    second = upload(client)
    assert client.delete("/api/session").status_code == 204
    assert client.get("/api/datasets/" + second["id"]).status_code == 401
    assert len(client.app.state.store.datasets) == 0


def test_production_cookie_and_origin_checks():
    settings = Settings(production=True, origin="https://dataforge.example")
    with TestClient(create_app(settings), base_url="https://testserver") as client:
        response = client.get("/api/session")
        cookie = response.headers["set-cookie"]
        assert "Secure" in cookie and "HttpOnly" in cookie and "SameSite=lax" in cookie
        assert client.post("/api/samples/sales").status_code == 403
        assert (
            client.post(
                "/api/samples/sales", headers={"Origin": "https://evil.example"}
            ).status_code
            == 403
        )
        assert (
            client.post("/api/samples/sales", headers={"Origin": settings.origin}).status_code
            == 201
        )


def test_absolute_expiry_deletes_private_frames():
    now = [0.0]
    store = WorkspaceStore(Settings(ttl=60), clock=lambda: now[0])
    token, session, _ = store.start(None)
    frame = pd.DataFrame({"x": [1]})
    d = store.add(token, "x.csv", frame, {})
    now[0] = 61
    store.prune()
    assert store.datasets == {}
    assert store.sessions == {}
    with pytest.raises(DataError):
        store.details(token, d["id"])


def test_memory_and_dataset_count_limits():
    small = WorkspaceStore(Settings(max_bytes=1))
    token, _, _ = small.start(None)
    with pytest.raises(DataError, match="memory"):
        small.add(token, "x.csv", pd.DataFrame({"x": [1]}), {})
    store = WorkspaceStore(Settings())
    token, _, _ = store.start(None)
    for n in range(3):
        store.add(token, str(n) + ".csv", pd.DataFrame({"x": [1]}), {})
    with pytest.raises(DataError, match="three datasets"):
        store.add(token, "four.csv", pd.DataFrame({"x": [1]}), {})


def test_concurrent_edits_allow_only_one_revision():
    store = WorkspaceStore(Settings())
    token, _, _ = store.start(None)
    d = store.add(token, "x.csv", pd.DataFrame({"x": pd.Series([" x "], dtype="string")}), {})

    def edit():
        try:
            store.transform(token, d["id"], 1, Operation(operation="trim_whitespace"))
            return 200
        except DataError as e:
            return e.status

    with ThreadPoolExecutor(max_workers=2) as pool:
        assert sorted(pool.map(lambda _: edit(), range(2))) == [200, 409]


def test_twenty_step_cap_and_rate_limits():
    store = WorkspaceStore(Settings())
    token, _, _ = store.start(None)
    d = store.add(token, "x.csv", pd.DataFrame({"x": pd.Series([" x "], dtype="string")}), {})
    for n in range(20):
        d = store.transform(token, d["id"], d["revision"], Operation(operation="trim_whitespace"))
    with pytest.raises(DataError, match="20 transformations"):
        store.transform(token, d["id"], d["revision"], Operation(operation="trim_whitespace"))
    with TestClient(create_app(Settings(request_limit=2))) as client:
        assert client.get("/api/health").status_code == 200
        assert client.get("/api/health").status_code == 200
        assert client.get("/api/health").status_code == 429


def test_optional_static_frontend_does_not_expose_source(tmp_path, monkeypatch):
    import app.main as main
    from fastapi.testclient import TestClient

    monkeypatch.setenv("SERVE_WEB", "true")
    monkeypatch.setattr(main, "ROOT", tmp_path)
    (tmp_path / "dist" / "assets").mkdir(parents=True)
    (tmp_path / "dist" / "index.html").write_text("<h1>DataForge</h1>")
    (tmp_path / "dist" / "assets" / "app.js").write_text("console.log('app')")
    (tmp_path / "private.txt").write_text("not public")
    with TestClient(main.create_app()) as client:
        assert client.get("/").text == "<h1>DataForge</h1>"
        assert client.get("/workspace").text == "<h1>DataForge</h1>"
        assert client.get("/private.txt").text == "<h1>DataForge</h1>"
        assert client.get("/assets/app.js").status_code == 200
        assert client.get("/api/not-a-route").status_code == 404
