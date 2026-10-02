from io import BytesIO
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
import json
import pandas as pd
from pandas.testing import assert_frame_equal
import pytest
from openpyxl import Workbook, load_workbook
from app.errors import DataError
from app.parsing import parse_file, kind, inspect_workbook
from app.profiling import profile, preview, chart_data, duplicate_candidates
from app.transformations import apply_transform, convert
from app.exporting import export_csv, export_xlsx
from app.schemas import Operation

FIXTURES = Path(__file__).parent / "fixtures"


def dirty():
    return parse_file("dirty.csv", (FIXTURES / "dirty.csv").read_bytes())[0]


def op(frame, operation, **kwargs):
    return apply_transform(frame, Operation(operation=operation, **kwargs))[0]


def workbook(rows, sheet="Data"):
    book = Workbook()
    ws = book.active
    ws.title = sheet
    for row in rows:
        ws.append(row)
    stream = BytesIO()
    book.save(stream)
    book.close()
    return stream.getvalue()


def test_inference_preserves_identifiers_missing_and_source():
    frame = dirty()
    assert list(frame["Customer ID"]) == ["001", "002", "003", "003"]
    assert kind(frame["Revenue"]) == "integer"
    assert frame["Revenue"].isna().sum() == 1
    assert frame.loc[0, "Region"] == " North "


def test_profile_numeric_statistics_and_duplicates():
    result = profile(dirty())
    assert result["rows"] == 4
    assert result["columns"] == 4
    assert result["missingCells"] == 3
    assert result["duplicateRows"] == 1
    revenue = next(c for c in result["columnProfiles"] if c["name"] == "Revenue")
    assert revenue["stats"]["mean"] == pytest.approx(700 / 3)
    assert revenue["stats"]["median"] == 300
    assert revenue["stats"]["min"] == 100
    json.dumps(result, allow_nan=False)


def test_remove_duplicates_is_pure():
    frame = dirty()
    before = frame.copy(deep=True)
    result = op(frame, "remove_duplicates")
    assert len(result) == 3
    assert_frame_equal(frame, before)


def test_drop_missing_selected_columns():
    assert len(op(dirty(), "drop_missing", columns=["Revenue"])) == 3
    assert len(op(dirty(), "drop_missing")) == 1


@pytest.mark.parametrize("method,expected", [("fill_mean", 700 / 3), ("fill_median", 300)])
def test_fill_numeric(method, expected):
    result = op(dirty(), method, column="Revenue")
    assert result.loc[1, "Revenue"] == pytest.approx(expected)
    assert result["Revenue"].isna().sum() == 0


def test_fractional_mean_promotes_integer_without_truncation():
    frame = parse_file("x.csv", b'x\n1\n2\n""\n')[0]
    result = op(frame, "fill_mean", column="x")
    assert result.loc[2, "x"] == 1.5
    assert kind(result["x"]) == "decimal"


def test_numeric_fill_rejects_empty_and_text_columns():
    with pytest.raises(DataError):
        op(dirty(), "fill_mean", column="Region")
    frame = dirty()
    frame["Revenue"] = pd.Series([pd.NA] * 4, dtype="Float64")
    with pytest.raises(DataError):
        op(frame, "fill_mean", column="Revenue")


def test_categorical_fill_retains_other_values():
    result = op(dirty(), "fill_category", column="Region", value="Unknown")
    assert result.loc[2, "Region"] == "Unknown"
    assert result.loc[0, "Region"] == " North "


@pytest.mark.parametrize(
    "operation,kwargs",
    [
        ("rename_column", {"column": "Revenue", "value": "Region"}),
        ("rename_column", {"column": "Revenue", "value": " "}),
        ("drop_columns", {"columns": ["Customer ID", "Customer Name", "Revenue", "Region"]}),
        ("drop_missing", {"columns": ["missing"]}),
        ("fill_category", {"column": "Region", "value": " "}),
    ],
)
def test_invalid_operations_fail_atomically(operation, kwargs):
    frame = dirty()
    before = frame.copy(deep=True)
    with pytest.raises(DataError):
        op(frame, operation, **kwargs)
    assert_frame_equal(frame, before)


def test_rename_and_remove_columns():
    renamed = op(dirty(), "rename_column", column="Revenue", value="Sales")
    assert "Sales" in renamed
    assert list(op(renamed, "drop_columns", columns=["Region", "Customer Name"]).columns) == [
        "Customer ID",
        "Sales",
    ]


def test_normalize_names_resolves_collisions_and_unicode():
    frame = pd.DataFrame([[1, 2, 3, 4]], columns=[" A B ", "A-B", "Ștefan Name", "!!!"])
    result = op(frame, "normalize_names")
    assert list(result.columns) == ["a_b", "a_b_2", "stefan_name", "column"]


def test_trim_whitespace_does_not_change_numeric_or_missing():
    result = op(dirty(), "trim_whitespace")
    assert result.loc[0, "Region"] == "North"
    assert pd.isna(result.loc[2, "Region"])
    assert kind(result["Revenue"]) == "integer"


@pytest.mark.parametrize(
    "target,values,expected",
    [
        ("integer", ["12", " 14 ", None], 12),
        ("decimal", ["1.25", "2", None], 1.25),
        ("boolean", ["yes", "false", None], True),
        ("date", ["2026-01-12", "2026-02-01", None], pd.Timestamp("2026-01-12")),
        ("text", [1, 2, None], "1.0"),
    ],
)
def test_compatible_conversion(target, values, expected):
    result = convert(pd.Series(values), target)
    assert result.iloc[0] == expected
    assert pd.isna(result.iloc[-1])


@pytest.mark.parametrize(
    "target,values",
    [
        ("integer", ["1.2", "2"]),
        ("decimal", ["inf", "2"]),
        ("decimal", ["9007199254740992", "2"]),
        ("boolean", ["perhaps", "yes"]),
        ("date", ["2026-02-30", "2026-01-01"]),
        ("date", ["01/02/2026", "2026-01-01"]),
    ],
)
def test_incompatible_conversion_rejected(target, values):
    with pytest.raises(DataError):
        convert(pd.Series(values, dtype="string"), target)


def test_duplicate_candidates_are_advisory_and_non_mutating():
    frame = pd.DataFrame(
        {"Name": [" Ava Chen ", "ava chen", "Noah Smith", "Noah Smyth", "Maya"]}, dtype="string"
    )
    before = frame.copy(deep=True)
    result = duplicate_candidates(frame, "Name")
    assert len(result["candidates"]) >= 2
    assert_frame_equal(frame, before)


def test_histogram_counts_and_scatter_filter_missing():
    frame = pd.DataFrame(
        {
            "x": pd.Series([1, 2, 3, pd.NA], dtype="Float64"),
            "y": pd.Series([3, 4, pd.NA, 5], dtype="Float64"),
        }
    )
    hist = chart_data(frame, "histogram", "x")
    assert sum(v["y"] for v in hist["data"]) == 3
    assert chart_data(frame, "scatter", "x", "y")["includedRows"] == 2


def test_bar_and_line_are_bounded_and_sorted():
    frame = pd.DataFrame({"label": ["B", "A", "C"], "amount": [20, 10, 30]})
    assert chart_data(frame, "bar", "label")["includedRows"] == 3
    numeric = pd.DataFrame({"x": [3, 1, 2], "y": [30, 10, 20]})
    assert chart_data(numeric, "line", "x", "y")["data"][0] == {"x": 1, "y": 10}
    large = pd.DataFrame({"x": range(500), "y": range(500)})
    assert len(chart_data(large, "scatter", "x", "y")["data"]) == 400
    assert chart_data(large, "line", "x", "y")["limited"]


def test_chart_invalid_axes_and_empty_data():
    with pytest.raises(DataError):
        chart_data(dirty(), "histogram", "Region")
    with pytest.raises(DataError):
        chart_data(dirty(), "line", "Revenue", "Revenue")
    empty = dirty().iloc[:0]
    assert profile(empty)["rows"] == 0
    assert chart_data(empty, "histogram", "Revenue")["data"] == []


def test_preview_paginates_and_serializes_missing():
    result = preview(dirty(), 1, 2)
    assert len(result["rows"]) == 2
    assert result["rows"][0][2] is None
    assert result["offset"] == 1


def test_csv_bom_and_formula_injection_defense():
    frame = pd.DataFrame(
        {"=header": ["=SUM(A1)", " +cmd", "Ștefan", "a,b"], "number": [-1, 2, 3, 4]}
    )
    data = export_csv(frame).decode("utf-8-sig")
    assert "'=header" in data
    assert "'=SUM(A1)" in data
    assert "' +cmd" in data
    assert "Ștefan" in data
    assert "-1" in data


def test_excel_exports_literal_strings_and_missing_values():
    frame = pd.DataFrame({"Name": ["=1+1", "Ștefan", None], "Number": [1.5, -2, None]})
    data = export_xlsx(frame)
    book = load_workbook(BytesIO(data), data_only=False)
    ws = book.active
    assert ws["A2"].value == "=1+1"
    assert ws["A2"].data_type == "s"
    assert ws["A3"].value == "Ștefan"
    assert ws["A4"].value is None
    assert ws.freeze_panes == "A2"
    book.close()


def test_csv_delimiters_and_utf8_bom():
    frame, meta = parse_file("data.csv", "\ufeffName;Value\nȘtefan;12\n".encode())
    assert len(frame.columns) == 2
    assert meta["delimiter"] == ";"
    assert frame.iloc[0, 0] == "Ștefan"


@pytest.mark.parametrize(
    "name", ["malformed-row.csv", "duplicate-headers.csv", "unclosed-quote.csv"]
)
def test_malformed_csv_fixtures(name):
    with pytest.raises(DataError):
        parse_file(name, (FIXTURES / name).read_bytes())


@pytest.mark.parametrize(
    "name,data",
    [
        ("x.csv", b""),
        ("x.csv", b"\xff\xfe"),
        ("x.csv", b"a\x00\nb\n"),
        ("x.csv", b"a,b\n"),
        ("x.xlsm", b"abc"),
        ("x.xlsx", b"not zip"),
        ("x.csv", b"a\n" + b"x" * (5 * 1024 * 1024)),
    ],
)
def test_invalid_uploads(name, data):
    with pytest.raises(DataError):
        parse_file(name, data)


def test_excel_sheet_selection_and_inference():
    data = workbook([["ID", "Quantity"], ["001", 4], ["002", None]])
    frame, meta = parse_file("x.xlsx", data)
    assert meta["sheet"] == "Data"
    assert frame.loc[0, "ID"] == "001"
    assert kind(frame["Quantity"]) == "integer"
    with pytest.raises(DataError):
        parse_file("x.xlsx", data, sheet="Missing")


def test_excel_formula_rejected():
    with pytest.raises(DataError, match="Formulas"):
        parse_file("x.xlsx", workbook([["Value"], ["=1+1"]]))


def test_excel_macro_rejected_without_execution():
    data = workbook([["Name"], ["Ava"]])
    out = BytesIO()
    with ZipFile(BytesIO(data)) as source, ZipFile(out, "w", ZIP_DEFLATED) as archive:
        for item in source.infolist():
            archive.writestr(item.filename, source.read(item))
        archive.writestr("xl/vbaProject.bin", b"not executable")
    with pytest.raises(DataError, match="Macros"):
        inspect_workbook(out.getvalue())


def test_zip_expansion_bomb_rejected():
    out = BytesIO()
    with ZipFile(out, "w", ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", b"x")
        z.writestr("xl/workbook.xml", b"x")
        z.writestr("xl/worksheets/sheet1.xml", b"A" * 1000000)
    with pytest.raises(DataError, match="compression"):
        inspect_workbook(out.getvalue())


def test_excel_sparse_extreme_coordinates_rejected():
    book = Workbook()
    ws = book.active
    ws.append(["Name"])
    ws["XFD99999"] = "No"
    out = BytesIO()
    book.save(out)
    book.close()
    with pytest.raises(DataError, match="coordinates"):
        parse_file("x.xlsx", out.getvalue())
